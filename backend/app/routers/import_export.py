import csv
import io
import shutil
from datetime import datetime
from pathlib import Path

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session, joinedload

from app.config import settings
from app.database import get_db
from app.models import Item, Location, Stock, TrackMode
from app.schemas import ItemCreate
from app.services import items as item_service
from app.services.barcode import next_barcode

router = APIRouter(prefix="/api/data", tags=["data"])

IMPORT_COLUMNS = [
    "name_zh",
    "name_en",
    "program",
    "track_mode",
    "category",
    "spec",
    "unit",
    "min_stock",
    "location_code",
    "quantity_initial",
    "note",
    "barcode",
]


@router.get("/export/csv")
def export_items_csv(db: Session = Depends(get_db)):
    rows = (
        db.query(Item)
        .options(joinedload(Item.stock), joinedload(Item.location))
        .filter(Item.active.is_(True))
        .order_by(Item.barcode)
        .all()
    )
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(
        [
            "barcode",
            "name_zh",
            "name_en",
            "program",
            "track_mode",
            "category",
            "spec",
            "unit",
            "min_stock",
            "location_code",
            "quantity",
            "note",
        ]
    )
    for item in rows:
        writer.writerow(
            [
                item.barcode,
                item.name_zh,
                item.name_en,
                item.program,
                item.track_mode,
                item.category,
                item.spec,
                item.unit,
                item.min_stock,
                item.location.code if item.location else "",
                item.stock.quantity if item.stock else 0,
                item.note,
            ]
        )
    buf.seek(0)
    return StreamingResponse(
        iter([buf.getvalue().encode("utf-8-sig")]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename=items_export.csv"},
    )


@router.get("/export/print-csv")
def export_print_csv(db: Session = Depends(get_db)):
    """CSV for 璞趣 PQPrint batch import."""
    rows = (
        db.query(Item)
        .filter(Item.active.is_(True))
        .order_by(Item.barcode)
        .all()
    )
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["barcode", "name_zh", "name_en", "program", "track_mode"])
    for item in rows:
        writer.writerow(
            [item.barcode, item.name_zh, item.name_en, item.program, item.track_mode]
        )
    buf.seek(0)
    return StreamingResponse(
        iter([buf.getvalue().encode("utf-8-sig")]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename=labels_pqprint.csv"},
    )


@router.get("/export/template.xlsx")
def download_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "items"
    ws.append(IMPORT_COLUMNS)
    ws.append(
        [
            "NEO 电机",
            "NEO Motor",
            "FRC",
            "SNP",
            "BEAR",
            "2-1/16 in",
            "pcs",
            2,
            "A-01-01",
            1,
            "",
            "",
        ]
    )
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": "attachment; filename=items_import_template.xlsx"
        },
    )


def _resolve_location(db: Session, code: str | None) -> int | None:
    if not code or not str(code).strip():
        return None
    code = str(code).strip()
    loc = db.query(Location).filter(Location.code == code).first()
    if not loc:
        loc = Location(code=code, name_zh=code, name_en=code)
        db.add(loc)
        db.flush()
    return loc.id


@router.post("/import")
async def import_items(file: UploadFile = File(...), db: Session = Depends(get_db)):
    name = file.filename or ""
    content = await file.read()
    rows: list[dict] = []

    if name.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(content), read_only=True)
        ws = wb.active
        headers = [str(c.value or "").strip() for c in next(ws.iter_rows(max_row=1))]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            rows.append(dict(zip(headers, row)))
    elif name.endswith(".csv"):
        text = content.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        rows = list(reader)
    else:
        raise HTTPException(400, "Upload .xlsx or .csv")

    created = 0
    errors: list[str] = []
    for i, row in enumerate(rows, start=2):
        try:
            name_zh = str(row.get("name_zh") or "").strip()
            name_en = str(row.get("name_en") or "").strip()
            if not name_zh and not name_en:
                continue
            program = str(row.get("program") or "BOTH").strip().upper()
            track_mode = str(row.get("track_mode") or "BLK").strip().upper()
            if program not in ("FRC", "FTC", "BOTH"):
                program = "BOTH"
            if track_mode not in ("SNP", "BLK"):
                track_mode = "BLK"

            barcode = row.get("barcode")
            if barcode:
                barcode = str(barcode).strip()
            else:
                barcode = next_barcode(db, program, track_mode)

            if db.query(Item).filter(Item.barcode == barcode).first():
                errors.append(f"Row {i}: duplicate barcode {barcode}")
                continue

            loc_id = _resolve_location(
                db, row.get("location_code") if row.get("location_code") else None
            )
            qty = int(float(row.get("quantity_initial") or 0))
            min_stock = int(float(row.get("min_stock") or 0))

            item = Item(
                barcode=barcode,
                program=program,
                track_mode=track_mode,
                name_zh=name_zh or name_en,
                name_en=name_en or name_zh,
                category=str(row.get("category") or "MISC").strip(),
                spec=str(row.get("spec") or "").strip(),
                unit=str(row.get("unit") or "pcs").strip(),
                min_stock=min_stock,
                location_id=loc_id,
                note=str(row.get("note") or "").strip(),
            )
            db.add(item)
            db.flush()
            db.add(Stock(item_id=item.id, quantity=max(0, qty)))
            created += 1
        except Exception as e:
            errors.append(f"Row {i}: {e}")

    db.commit()
    return {"created": created, "errors": errors}


@router.post("/backup")
def backup_database():
    src = settings.data_dir / "inventory.db"
    if not src.exists():
        raise HTTPException(404, "Database not found")
    backups = settings.data_dir / "backups"
    backups.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    dest = backups / f"inventory_{stamp}.db"
    shutil.copy2(src, dest)
    return {"path": str(dest), "filename": dest.name}
